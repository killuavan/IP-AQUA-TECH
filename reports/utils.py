"""Report generation utilities – PDF (ReportLab) and Excel (openpyxl)."""
import io
from datetime import datetime, time
from decimal import Decimal
from types import SimpleNamespace

from django.utils import timezone

from iot.firebase_client import FirebaseNotConfigured, get_firestore


def gather_report_data(start_date, end_date):
    """Aggregate report data from the Firestore RTDB mirror history."""
    try:
        history = list(get_firestore().collection('rtdb_mirror_history').stream())
    except (FirebaseNotConfigured, Exception) as exc:
        raise RuntimeError('Firestore is unavailable; report data was not loaded.') from exc

    start = timezone.make_aware(datetime.combine(start_date, time.min))
    end = timezone.make_aware(datetime.combine(end_date, time.max))
    grouped = {'water_quality': [], 'solar': [], 'battery': [], 'tank': [], 'pump': []}
    for document in history:
        value = document.to_dict() or {}
        path = value.get('path', '')
        label = path.removeprefix('sensors/').removesuffix('/latest')
        received_at = value.get('received_at')
        if label not in grouped or not received_at or not (start <= received_at <= end):
            continue
        grouped[label].append(value.get('data') or {})

    def number(payload, key, default=0):
        value = payload.get(key, default)
        try:
            return float(value)
        except (TypeError, ValueError):
            return default

    water_quality = []
    for payload in grouped['water_quality']:
        water_quality.append(SimpleNamespace(
            recorded_at=timezone.now(),
            ph=number(payload, 'ph'),
            tds=number(payload, 'tds'),
            turbidity=number(payload, 'turbidity'),
            temperature=number(payload, 'temperature'),
            flow_rate=number(payload, 'flow_rate'),
            water_pressure=number(payload, 'water_pressure'),
            water_level=number(payload, 'water_level'),
            status=payload.get('status', payload.get('warning_level', 'NORMAL')),
        ))

    def average(records, field):
        values = [getattr(record, field) for record in records]
        return sum(values) / len(values) if values else None

    solar = grouped['solar']
    solar_power = [number(item, 'power_output') for item in solar]

    return {
        'start_date': start_date,
        'end_date': end_date,
        'generated_at': timezone.now(),
        'water_quality': {
            'count': len(water_quality),
            'avg_ph': average(water_quality, 'ph'),
            'avg_tds': average(water_quality, 'tds'),
            'avg_turbidity': average(water_quality, 'turbidity'),
            'avg_temperature': average(water_quality, 'temperature'),
            'critical': sum(record.status == 'CRITICAL' for record in water_quality),
            'warning': sum(record.status == 'WARNING' for record in water_quality),
            'records': water_quality[:100],
        },
        'production': {
            'total_liters': 0,
            'avg_daily': 0,
            'days': 0,
            'records': [],
        },
        'solar': {
            'avg_power': sum(solar_power) / len(solar_power) if solar_power else 0,
            'readings': len(solar),
            'energy_kwh': sum(number(item, 'energy_generated_kwh') for item in solar),
            'records': [],
        },
        'battery': grouped['battery'][-1] if grouped['battery'] else None,
        'tank': grouped['tank'][-1] if grouped['tank'] else None,
        'pump': grouped['pump'][-1] if grouped['pump'] else None,
        'filters': [],
        'maintenance': {
            'total': 0, 'completed': 0, 'pending': 0, 'records': [],
        },
        'alerts': {
            'total': 0, 'critical': 0, 'warning': 0, 'records': [],
        },
    }


def _fmt(val, digits=2):
    if val is None:
        return 'N/A'
    if isinstance(val, Decimal):
        return f'{float(val):.{digits}f}'
    try:
        return f'{float(val):.{digits}f}'
    except (TypeError, ValueError):
        return str(val)


def generate_pdf(data, title):
    """Return PDF bytes for the report."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable,
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        rightMargin=1.5 * cm, leftMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name='AquaTitle', parent=styles['Heading1'],
        textColor=colors.HexColor('#0a4d68'), fontSize=18, spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        name='AquaHead', parent=styles['Heading2'],
        textColor=colors.HexColor('#088395'), fontSize=13, spaceBefore=12, spaceAfter=6,
    ))

    story = []
    story.append(Paragraph('IP Aqua Tech', styles['AquaTitle']))
    story.append(Paragraph(
        'Solar-Powered Drinking Water Filtration Monitoring System',
        styles['Normal'],
    ))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width='100%', color=colors.HexColor('#05bfdb'), thickness=2))
    story.append(Spacer(1, 8))
    story.append(Paragraph(title, styles['Heading2']))
    story.append(Paragraph(
        f'Period: {data["start_date"]} to {data["end_date"]}<br/>'
        f'Generated: {data["generated_at"]:%Y-%m-%d %H:%M}',
        styles['Normal'],
    ))
    story.append(Spacer(1, 12))

    # Water Quality
    wq = data['water_quality']
    story.append(Paragraph('Water Quality Summary', styles['AquaHead']))
    wq_table = Table([
        ['Metric', 'Value'],
        ['Readings', str(wq['count'])],
        ['Avg pH', _fmt(wq['avg_ph'])],
        ['Avg TDS (ppm)', _fmt(wq['avg_tds'])],
        ['Avg Turbidity (NTU)', _fmt(wq['avg_turbidity'])],
        ['Avg Temperature (°C)', _fmt(wq['avg_temperature'])],
        ['Critical Readings', str(wq['critical'])],
        ['Warning Readings', str(wq['warning'])],
    ], colWidths=[10 * cm, 6 * cm])
    wq_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a4d68')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#e8f7fa')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(wq_table)

    # Production
    prod = data['production']
    story.append(Paragraph('Water Production', styles['AquaHead']))
    prod_table = Table([
        ['Metric', 'Value'],
        ['Total Produced (L)', _fmt(prod['total_liters'])],
        ['Average Daily (L)', _fmt(prod['avg_daily'])],
        ['Days Recorded', str(prod['days'])],
    ], colWidths=[10 * cm, 6 * cm])
    prod_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a4d68')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#e8f7fa')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(prod_table)

    # Solar
    sol = data['solar']
    story.append(Paragraph('Solar Energy', styles['AquaHead']))
    sol_table = Table([
        ['Metric', 'Value'],
        ['Avg Power Output (W)', _fmt(sol['avg_power'])],
        ['Readings', str(sol['readings'])],
        ['Energy Generated (kWh)', _fmt(sol['energy_kwh'], 3)],
    ], colWidths=[10 * cm, 6 * cm])
    sol_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a4d68')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#e8f7fa')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(sol_table)

    # Maintenance & Alerts
    maint = data['maintenance']
    alerts = data['alerts']
    story.append(Paragraph('Maintenance & Alerts', styles['AquaHead']))
    ma_table = Table([
        ['Metric', 'Value'],
        ['Maintenance Tasks', str(maint['total'])],
        ['Completed', str(maint['completed'])],
        ['Pending', str(maint['pending'])],
        ['Total Alerts', str(alerts['total'])],
        ['Critical Alerts', str(alerts['critical'])],
        ['Warning Alerts', str(alerts['warning'])],
    ], colWidths=[10 * cm, 6 * cm])
    ma_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a4d68')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#e8f7fa')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(ma_table)

    story.append(Spacer(1, 20))
    story.append(Paragraph(
        '— End of Report — IP Aqua Tech Monitoring System —',
        styles['Normal'],
    ))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel(data, title):
    """Return Excel (xlsx) bytes for the report."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side

    wb = Workbook()

    header_fill = PatternFill('solid', fgColor='0A4D68')
    header_font = Font(bold=True, color='FFFFFF')
    aqua_fill = PatternFill('solid', fgColor='E8F7FA')
    thin = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )

    # Summary sheet
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = 'IP Aqua Tech'
    ws['A1'].font = Font(bold=True, size=16, color='0A4D68')
    ws['A2'] = title
    ws['A3'] = f'Period: {data["start_date"]} to {data["end_date"]}'
    ws['A4'] = f'Generated: {data["generated_at"]:%Y-%m-%d %H:%M}'

    rows = [
        ('', ''),
        ('WATER QUALITY', ''),
        ('Readings', data['water_quality']['count']),
        ('Avg pH', _fmt(data['water_quality']['avg_ph'])),
        ('Avg TDS', _fmt(data['water_quality']['avg_tds'])),
        ('Avg Turbidity', _fmt(data['water_quality']['avg_turbidity'])),
        ('Avg Temperature', _fmt(data['water_quality']['avg_temperature'])),
        ('Critical', data['water_quality']['critical']),
        ('Warning', data['water_quality']['warning']),
        ('', ''),
        ('PRODUCTION', ''),
        ('Total Liters', _fmt(data['production']['total_liters'])),
        ('Avg Daily', _fmt(data['production']['avg_daily'])),
        ('Days', data['production']['days']),
        ('', ''),
        ('SOLAR', ''),
        ('Avg Power (W)', _fmt(data['solar']['avg_power'])),
        ('Energy (kWh)', _fmt(data['solar']['energy_kwh'], 3)),
        ('', ''),
        ('MAINTENANCE', ''),
        ('Total Tasks', data['maintenance']['total']),
        ('Completed', data['maintenance']['completed']),
        ('Pending', data['maintenance']['pending']),
        ('', ''),
        ('ALERTS', ''),
        ('Total', data['alerts']['total']),
        ('Critical', data['alerts']['critical']),
        ('Warning', data['alerts']['warning']),
    ]
    for i, (k, v) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
        if k and v == '':
            ws.cell(row=i, column=1).font = Font(bold=True, color='088395')

    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 18

    # Production detail
    ws2 = wb.create_sheet('Production')
    headers = ['Date', 'Daily Production (L)', 'Target (L)', 'Notes']
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
    for row_i, rec in enumerate(data['production']['records'], 2):
        ws2.cell(row=row_i, column=1, value=str(rec.date))
        ws2.cell(row=row_i, column=2, value=float(rec.daily_production_liters))
        ws2.cell(row=row_i, column=3, value=float(rec.target_liters))
        ws2.cell(row=row_i, column=4, value=rec.notes)
    for col in range(1, 5):
        ws2.column_dimensions[chr(64 + col)].width = 20

    # Water quality detail
    ws3 = wb.create_sheet('Water Quality')
    headers = ['Date/Time', 'pH', 'TDS', 'Turbidity', 'Temp', 'Flow', 'Pressure', 'Level', 'Status']
    for col, h in enumerate(headers, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
    for row_i, rec in enumerate(data['water_quality']['records'], 2):
        ws3.cell(row=row_i, column=1, value=rec.recorded_at.strftime('%Y-%m-%d %H:%M'))
        ws3.cell(row=row_i, column=2, value=float(rec.ph))
        ws3.cell(row=row_i, column=3, value=float(rec.tds))
        ws3.cell(row=row_i, column=4, value=float(rec.turbidity))
        ws3.cell(row=row_i, column=5, value=float(rec.temperature))
        ws3.cell(row=row_i, column=6, value=float(rec.flow_rate))
        ws3.cell(row=row_i, column=7, value=float(rec.water_pressure))
        ws3.cell(row=row_i, column=8, value=float(rec.water_level))
        ws3.cell(row=row_i, column=9, value=rec.status)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
